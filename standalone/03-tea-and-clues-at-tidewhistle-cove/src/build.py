"""Builds ../<prefix>.epub (EPUB 3, reflowable, text only, nav + NCX) and ../<prefix>.docx (for Kindle Create).
Run from src/: python3 build.py   (run cover.py first so the cover image can be embedded in the EPUB)"""
import os, zipfile, uuid, json, html, datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
import book

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
EPUB = os.path.join(ROOT, book.PREFIX + ".epub")
DOCX = os.path.join(ROOT, book.PREFIX + ".docx")
COVER = os.path.join(ROOT, book.PREFIX + "-cover.jpg")
BOOK_ID = "urn:uuid:" + str(uuid.uuid5(uuid.NAMESPACE_URL, "kdp-books/" + book.PREFIX))
e = html.escape

CSS = """body { font-family: serif; line-height: 1.5; margin: 0 4%; }
h1 { font-size: 1.5em; text-align: center; margin: 2em 0 1.2em; page-break-before: always; font-weight: bold; }
h2 { font-size: 1.2em; text-align: center; margin: 1.6em 0 1em; page-break-before: always; font-weight: bold; }
p { text-indent: 1.2em; margin: 0 0 0.5em; text-align: left; }
p.first, p.ask, p.think, p.center { text-indent: 0; }
p.ask { font-weight: bold; margin-top: 1.2em; }
p.think { font-style: italic; text-align: center; margin: 1.2em 0; }
p.center { text-align: center; }
p.title { text-indent: 0; text-align: center; font-size: 2em; font-weight: bold; margin-top: 3em; }
p.subtitle { text-indent: 0; text-align: center; font-size: 1.2em; font-style: italic; margin: 1em 0 2em; }
p.author { text-indent: 0; text-align: center; font-size: 1.3em; }
p.small { text-indent: 0; font-size: 0.85em; margin-bottom: 0.8em; }
nav ol { list-style-type: none; padding-left: 0; }
nav li { margin: 0.3em 0; }
"""

def xhtml(title, body):
    return f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="en" lang="en">
<head><meta charset="utf-8"/><title>{e(title)}</title><link rel="stylesheet" type="text/css" href="style.css"/></head>
<body>
{body}
</body>
</html>
"""

def section_body(sid, head, blocks):
    out = [f'<section epub:type="chapter" id="{sid}"><h1>{e(head)}</h1>']
    first = True
    for kind, text in blocks:
        if kind == "h2":
            out.append(f"<h2>{e(text)}</h2>"); first = True; continue
        cls = {"ask": "ask", "think": "think"}.get(kind, "first" if first else "")
        out.append(f'<p class="{cls}">{e(text)}</p>' if cls else f"<p>{e(text)}</p>")
        first = kind in ("ask", "think")
    out.append("</section>")
    return "\n".join(out)

def build_epub():
    secs = book.sections()
    files = []  # (filename, title, content)
    title_body = (f'<section epub:type="titlepage"><p class="title">{e(book.TITLE)}</p>'
                  f'<p class="subtitle">{e(book.BOOK_SUBTITLE)}</p><p class="author">{e(book.AUTHOR)}</p></section>')
    files.append(("title.xhtml", book.TITLE, title_body))
    copy_body = '<section epub:type="copyright-page">' + "".join(f'<p class="small">{e(t)}</p>' for t in book.COPYRIGHT) + "</section>"
    files.append(("copyright.xhtml", "Copyright", copy_body))
    toc_items = "\n".join(f'<li><a href="{sid}.xhtml">{e(h)}</a></li>' for sid, h, _ in secs)
    files.append(("contents.xhtml", "Contents", f'<section epub:type="toc"><h1>Contents</h1><ol>{toc_items}</ol></section>'))
    for sid, h, blocks in secs:
        files.append((f"{sid}.xhtml", h, section_body(sid, h, blocks)))
    nav = xhtml("Contents", f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>\n{toc_items}\n</ol></nav>\n'
                f'<nav epub:type="landmarks" hidden=""><ol><li><a epub:type="toc" href="contents.xhtml">Contents</a></li>'
                f'<li><a epub:type="bodymatter" href="{secs[0][0]}.xhtml">Start</a></li></ol></nav>')
    navpoints = "\n".join(f'<navPoint id="np{i+1}" playOrder="{i+1}"><navLabel><text>{e(h)}</text></navLabel><content src="{sid}.xhtml"/></navPoint>'
                          for i, (sid, h, _) in enumerate(secs))
    ncx = f"""<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head><meta name="dtb:uid" content="{BOOK_ID}"/>
<meta name="dtb:depth" content="1"/><meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head>
<docTitle><text>{e(book.TITLE)}</text></docTitle><navMap>
{navpoints}
</navMap></ncx>
"""
    has_cover = os.path.exists(COVER)
    manifest = ['<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
                '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
                '<item id="css" href="style.css" media-type="text/css"/>']
    if has_cover:
        manifest.append('<item id="cover-image" href="cover.jpg" media-type="image/jpeg" properties="cover-image"/>')
    spine = []
    for fn, _, _ in files:
        iid = fn.replace(".xhtml", "")
        manifest.append(f'<item id="{iid}" href="{fn}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="{iid}"/>')
    modified = "2026-10-03T00:00:00Z"
    opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="en">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="bookid">{BOOK_ID}</dc:identifier>
<dc:title>{e(book.TITLE)}</dc:title>
<dc:creator id="author">{e(book.AUTHOR)}</dc:creator>
<meta refines="#author" property="role" scheme="marc:relators">aut</meta>
<dc:language>en</dc:language>
<dc:publisher>{e(book.AUTHOR)}</dc:publisher>
<dc:date>2026-10-03</dc:date>
<dc:description>{e(book.SUBTITLE)}</dc:description>
<meta property="dcterms:modified">{modified}</meta>
{'<meta name="cover" content="cover-image"/>' if has_cover else ''}
</metadata>
<manifest>
{chr(10).join(manifest)}
</manifest>
<spine toc="ncx">
{chr(10).join(spine)}
</spine>
</package>
"""
    container = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>
<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>
"""
    fixed = (2026, 10, 3, 0, 0, 0)
    with zipfile.ZipFile(EPUB, "w") as z:
        def add(name, data, compress=zipfile.ZIP_DEFLATED):
            zi = zipfile.ZipInfo(name, fixed); zi.compress_type = compress
            z.writestr(zi, data)
        add("mimetype", "application/epub+zip", zipfile.ZIP_STORED)
        add("META-INF/container.xml", container)
        add("OEBPS/content.opf", opf)
        add("OEBPS/nav.xhtml", nav)
        add("OEBPS/toc.ncx", ncx)
        add("OEBPS/style.css", CSS)
        if has_cover: add("OEBPS/cover.jpg", open(COVER, "rb").read())
        for fn, t, body in files:
            add("OEBPS/" + fn, xhtml(t, body))
    return os.path.getsize(EPUB)

def build_docx():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Georgia"; st.font.size = Pt(12)
    def center(text, size, bold=False, italic=False):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        return p
    center(book.TITLE, 26, bold=True)
    center(book.BOOK_SUBTITLE, 15, italic=True)
    center(book.AUTHOR, 16)
    d.add_page_break()
    for t in book.COPYRIGHT:
        p = d.add_paragraph(t); p.runs[0].font.size = Pt(10)
    for sid, h, blocks in book.sections():
        d.add_page_break()
        d.add_heading(h, level=1)
        for kind, text in blocks:
            if kind == "h2":
                d.add_page_break(); d.add_heading(text, level=2); continue
            p = d.add_paragraph()
            r = p.add_run(text)
            if kind == "ask": r.bold = True
            if kind == "think": r.italic = True; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.core_properties.title = book.TITLE
    d.core_properties.author = book.AUTHOR
    d.save(DOCX)
    return os.path.getsize(DOCX)

if __name__ == "__main__":
    eb = build_epub(); db = build_docx()
    words = len(book.all_text().split())
    price = 3.99
    mb = round(eb / 1024) / 1024  # size rounded to nearest KB, in MB
    import math
    fee = max(0.01, round(math.ceil(eb / 1024) / 1024 * 0.15 + 1e-9, 2))
    roy = round(0.70 * price - fee, 2)
    info = {"title": book.TITLE, "subtitle": book.SUBTITLE, "author": book.AUTHOR, "format": "Kindle ebook (EPUB 3, reflowable, text only) + DOCX for Kindle Create",
            "audio_first": True, "cases": len(book.CASES), "words": words, "epub_bytes": eb, "docx_bytes": db,
            "price_usd": price, "royalty_option": "70%", "delivery_fee_usd_per_mb": 0.15, "delivery_fee_usd_est": fee,
            "royalty_usd_est": roy, "royalty_math": f"0.70 x {price} - {fee:.2f} = {roy:.2f}",
            "files": {"epub": os.path.basename(EPUB), "docx": os.path.basename(DOCX), "cover": os.path.basename(COVER)}}
    json.dump(info, open(os.path.join(ROOT, "build-info.json"), "w"), indent=2)
    print(info)
